"""ຊັ້ນຖານຂໍ້ມູນ SQLite + ສູດຄິດໄລ່ (ຕາມສູດໃນ Excel ຂອງຜູ້ໃຊ້)."""
import hashlib
import hmac
import os
import re
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd

DB_PATH = Path(__file__).with_name("fund.db")
RATE_DEFAULT = 0.02  # ດອກເບ້ຍເງິນກູ້ 2%/ເດືອນ  (I = H * 2%)

# ຄໍລຳທີ່ຜູ້ໃຊ້ປ້ອນເອງ  (ຊື່ຄໍລຳ -> ຊື່ຫົວຂໍ້ພາສາລາວ)
INPUT_COLS = {
    "old_dep": "ເງິນຝາກທັງໝົດໃນເດືອນເກົ່າ",
    "dep": "ເງິນຝາກປະຈຳເດືອນນີ້",
    "wd": "ຖອນເງິນຝາກ",
    "loan_total": "ເງິນກູ້ຢືມທັງໝົດໃນເດືອນນີ້",
    "accrued_prev": "ດອກຄ້າງຈ່າຍແຕ່ເດືອນກ່ອນ",
    "interest_actual": "ດອກເບ້ຍເງິນກູ້ໃນຕົວຈິງ",
    "repay": "ສົ່ງຄືນເງິນກູ້",
    "new_loan": "ເງິນກູ້ຢືມໃໝ່ໃນເດືອນນີ້",
    "app_fee": "ໃບຄຳຮ້ອງ",
    "booklet_fee": "ເງິນຄ່າປຶ້ມນ້ອຍ",
}
CALC_COLS = {
    "dep_next": "ເງິນຝາກຍ້າຍໄປເດືອນໜ້າ",       # G = D + E - F
    "interest_due": "ດອກເບ້ຍຕາມເງິນກູ້",        # I = H * rate
    "accrued_next": "ດອກຄ້າງໄປເດືອນໜ້າ",       # N = I + J - K
    "loan_next": "ເງິນກູ້ຄ້າງໄປເດືອນໜ້າ",       # O = M + H - L
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS members(
    book_no INTEGER PRIMARY KEY, name TEXT NOT NULL, active INTEGER DEFAULT 1);
CREATE TABLE IF NOT EXISTS monthly(
    period TEXT NOT NULL, book_no INTEGER NOT NULL, tx_date TEXT,
    old_dep REAL DEFAULT 0, dep REAL DEFAULT 0, wd REAL DEFAULT 0,
    loan_total REAL DEFAULT 0, accrued_prev REAL DEFAULT 0,
    interest_actual REAL DEFAULT 0, repay REAL DEFAULT 0, new_loan REAL DEFAULT 0,
    app_fee REAL DEFAULT 0, booklet_fee REAL DEFAULT 0, closed INTEGER DEFAULT 0,
    PRIMARY KEY(period, book_no));
CREATE TABLE IF NOT EXISTS extras(
    id INTEGER PRIMARY KEY AUTOINCREMENT, period TEXT, label TEXT, amount REAL);
CREATE TABLE IF NOT EXISTS dividend(
    year INTEGER, book_no INTEGER, tx_date TEXT,
    capital REAL DEFAULT 0, div_interest REAL DEFAULT 0,
    wd_interest REAL DEFAULT 0, wd_capital REAL DEFAULT 0, receiver TEXT,
    PRIMARY KEY(year, book_no));
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
"""


def conn():
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    return c


def init():
    with conn() as c:
        c.executescript(SCHEMA)
        c.executescript(SCHEMA2)
        cols = [r[1] for r in c.execute("PRAGMA table_info(users)")]
        for name, ddl in (("range_from", "INTEGER"), ("range_to", "INTEGER"),
                          ("view_all", "INTEGER DEFAULT 1")):
            if name not in cols:
                c.execute(f"ALTER TABLE users ADD COLUMN {name} {ddl}")


def get_rate() -> float:
    with conn() as c:
        r = c.execute("SELECT value FROM settings WHERE key='rate'").fetchone()
    return float(r["value"]) if r else RATE_DEFAULT


def set_rate(v: float):
    with conn() as c:
        c.execute("INSERT OR REPLACE INTO settings VALUES('rate',?)", (str(v),))


# ---------- ສູດຄິດໄລ່ ----------
def add_calc(df: pd.DataFrame, rate: float | None = None) -> pd.DataFrame:
    rate = get_rate() if rate is None else rate
    df = df.copy()
    for k in INPUT_COLS:
        df[k] = pd.to_numeric(df[k], errors="coerce").fillna(0.0)
    df["dep_next"] = df.old_dep + df.dep - df.wd
    df["interest_due"] = df.loan_total * rate
    df["accrued_next"] = df.interest_due + df.accrued_prev - df.interest_actual
    df["loan_next"] = df.new_loan + df.loan_total - df.repay
    return df


def summary(df: pd.DataFrame, extra_loans: float = 0.0) -> dict:
    """ຕາມແຖວສະຫຼຸບໃນ Excel: ລາຍຮັບ = E+K+L+P+Q ; ລາຍຈ່າຍ = F+M"""
    income = df.dep.sum() + df.interest_actual.sum() + df.repay.sum() \
        + df.app_fee.sum() + df.booklet_fee.sum()
    expense = df.wd.sum() + df.new_loan.sum()
    return {
        "income": income, "expense": expense, "balance": income - expense,
        "old_dep": df.old_dep.sum(), "dep": df.dep.sum(), "wd": df.wd.sum(),
        "dep_next": df.dep_next.sum(), "loan_total": df.loan_total.sum(),
        "loan_next": df.loan_next.sum(), "accrued_next": df.accrued_next.sum(),
        "interest_actual": df.interest_actual.sum(),
        "loan_with_extra": df.loan_total.sum() + extra_loans,
    }


# ---------- ງວດ (period = 'YYYY-MM') ----------
def periods() -> list[str]:
    with conn() as c:
        return [r[0] for r in c.execute(
            "SELECT DISTINCT period FROM monthly ORDER BY period")]


def next_period(p: str) -> str:
    y, m = map(int, p.split("-"))
    return f"{y + (m == 12)}-{(m % 12) + 1:02d}"


def load_month(period: str) -> pd.DataFrame:
    with conn() as c:
        df = pd.read_sql_query(
            """SELECT m.book_no, b.name, m.tx_date, m.old_dep, m.dep, m.wd,
                      m.loan_total, m.accrued_prev, m.interest_actual, m.repay,
                      m.new_loan, m.app_fee, m.booklet_fee, m.closed
               FROM monthly m LEFT JOIN members b USING(book_no)
               WHERE m.period=? ORDER BY m.book_no""", c, params=(period,))
    df["closed"] = df["closed"].fillna(0).astype(bool)
    return df


def save_month(period: str, df: pd.DataFrame):
    cols = ["tx_date"] + list(INPUT_COLS) + ["closed"]
    with conn() as c:
        for _, r in df.iterrows():
            vals = [None if pd.isna(r["tx_date"]) else str(r["tx_date"])[:10]]
            vals += [0.0 if pd.isna(r[k]) else float(r[k]) for k in INPUT_COLS]
            vals.append(int(bool(r["closed"])))
            c.execute(
                f"UPDATE monthly SET {', '.join(k + '=?' for k in cols)} "
                "WHERE period=? AND book_no=?", vals + [period, int(r["book_no"])])
            c.execute("UPDATE members SET name=? WHERE book_no=?",
                      (r["name"], int(r["book_no"])))


def rollover(period: str) -> str:
    """ສ້າງເດືອນໜ້າ: ຍອດຍົກມາຈາກເດືອນນີ້ອັດຕະໂນມັດ (ບໍ່ຕ້ອງພິມຄືນ)."""
    nxt = next_period(period)
    if nxt in periods():
        return nxt
    cur = add_calc(load_month(period))
    with conn() as c:
        for _, r in cur.iterrows():
            if r["closed"]:
                continue  # ບັນຊີທີ່ປິດແລ້ວບໍ່ຍົກໄປ
            c.execute(
                "INSERT INTO monthly(period,book_no,old_dep,loan_total,accrued_prev)"
                " VALUES(?,?,?,?,?)",
                (nxt, int(r.book_no), r.dep_next, r.loan_next, r.accrued_next))
    return nxt


def check_carry() -> pd.DataFrame:
    """ກວດວ່າ ຍອດເດືອນເກົ່າ (D) ຂອງເດືອນນີ້ ກົງກັບ ຍອດຍົກໄປເດືອນໜ້າ (G) ຂອງເດືອນກ່ອນບໍ່."""
    out, ps = [], periods()
    for a, b in zip(ps[:-1], ps[1:]):
        if next_period(a) != b:
            continue
        A, B = add_calc(load_month(a)).set_index("book_no"), load_month(b).set_index("book_no")
        for k in B.index.intersection(A.index):
            for src, dst, label in [("dep_next", "old_dep", "ເງິນຝາກ"),
                                    ("loan_next", "loan_total", "ເງິນກູ້")]:
                if abs(A.at[k, src] - B.at[k, dst]) > 0.5:
                    out.append({"ເດືອນ": b, "ເລກທີປຶ້ມ": k, "ລາຍການ": label,
                                "ຍອດຍົກຈາກເດືອນກ່ອນ": A.at[k, src],
                                "ຍອດທີ່ພິມໃນເດືອນນີ້": B.at[k, dst],
                                "ຜິດດ່ຽງ": B.at[k, dst] - A.at[k, src]})
    return pd.DataFrame(out)


# ---------- ສະມາຊິກ ----------
def members() -> pd.DataFrame:
    with conn() as c:
        return pd.read_sql_query("SELECT * FROM members ORDER BY book_no", c)


def add_member(book_no: int, name: str, period: str | None):
    with conn() as c:
        c.execute("INSERT INTO members(book_no,name) VALUES(?,?)", (book_no, name))
        if period:
            c.execute("INSERT OR IGNORE INTO monthly(period,book_no) VALUES(?,?)",
                      (period, book_no))


# ---------- ເງິນປັນຜົນທ້າຍປີ ----------
def load_dividend(year: int) -> pd.DataFrame:
    with conn() as c:
        df = pd.read_sql_query(
            """SELECT d.book_no, b.name, d.tx_date, d.capital, d.div_interest,
                      d.wd_interest, d.wd_capital, d.receiver
               FROM dividend d LEFT JOIN members b USING(book_no)
               WHERE d.year=? ORDER BY d.book_no""", c, params=(year,))
    df["total"] = df.capital + df.div_interest                     # F = D + E
    df["wd_total"] = df.wd_interest + df.wd_capital                # I = G + H
    df["carry_year"] = df.total - df.wd_total                      # J = F - I
    return df


def save_dividend(year: int, df: pd.DataFrame):
    with conn() as c:
        for _, r in df.iterrows():
            c.execute(
                "INSERT OR REPLACE INTO dividend VALUES(?,?,?,?,?,?,?,?)",
                (year, int(r.book_no),
                 None if pd.isna(r.tx_date) else str(r.tx_date)[:10],
                 float(r.capital or 0), float(r.div_interest or 0),
                 float(r.wd_interest or 0), float(r.wd_capital or 0),
                 None if pd.isna(r.receiver) else r.receiver))


# ---------- ນຳເຂົ້າຈາກ Excel ເກົ່າ ----------
def _num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def _sheet_period(name: str):
    nums = re.findall(r"\d+", name)
    if not nums:
        return None
    if len(nums) == 1 and len(nums[0]) >= 5:
        m, y = int(nums[0][:-4]), int(nums[0][-4:])
    elif len(nums) >= 2:
        m, y = int(nums[0]), int(nums[1])
    else:
        return None
    return f"{y}-{m:02d}" if 1 <= m <= 12 and y > 2000 else None


def import_excel(file, only_latest: bool = False) -> list[str]:
    """ນຳເຂົ້າທຸກຊີດປະຈຳເດືອນ (ຫຼືສະເພາະເດືອນຫຼ້າສຸດ). ຄືນລາຍຊື່ງວດທີ່ນຳເຂົ້າ."""
    from openpyxl import load_workbook
    wb = load_workbook(file, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        p = _sheet_period(ws.title)
        h = str(ws["D1"].value or "")
        if p and "ເງິນຝາກ" in h and "ເງິນກູ້" in str(ws["H1"].value or ""):
            sheets.append((p, ws))
    sheets.sort(key=lambda t: t[0])
    if only_latest and sheets:
        sheets = sheets[-1:]
    done = []
    with conn() as c:
        for p, ws in sheets:
            c.execute("DELETE FROM monthly WHERE period=?", (p,))
            for r in range(2, ws.max_row + 1):
                k = ws.cell(r, 1).value
                if not isinstance(k, (int, float)):
                    continue
                k = int(k)
                name = ws.cell(r, 2).value
                c.execute("INSERT OR IGNORE INTO members(book_no,name) VALUES(?,?)",
                          (k, str(name).strip() if name else f"ປຶ້ມເລກທີ {k}"))
                if name:
                    c.execute("UPDATE members SET name=? WHERE book_no=?",
                              (str(name).strip(), k))
                d = ws.cell(r, 3).value
                d = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else None
                c.execute(
                    "INSERT INTO monthly VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (p, k, d, _num(ws.cell(r, 4).value), _num(ws.cell(r, 5).value),
                     _num(ws.cell(r, 6).value), _num(ws.cell(r, 8).value),
                     _num(ws.cell(r, 10).value), _num(ws.cell(r, 11).value),
                     _num(ws.cell(r, 12).value), _num(ws.cell(r, 13).value),
                     _num(ws.cell(r, 16).value), _num(ws.cell(r, 17).value),
                     1 if ws.cell(r, 18).value and
                     str(ws.cell(r, 18).value).strip() and
                     "ປິດບັນຊີ" not in str(ws.cell(r, 18).value) else 0))
            done.append(p)
    return done


# =====================================================================
#  ລະບົບຜູ້ໃຊ້ / ສິດ / ບັນທຶກການເຄື່ອນໄຫວ
# =====================================================================
SCHEMA2 = """
CREATE TABLE IF NOT EXISTS users(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT UNIQUE NOT NULL, full_name TEXT NOT NULL, pw_hash TEXT NOT NULL,
    role TEXT NOT NULL DEFAULT 'user',        -- admin | user
    status TEXT NOT NULL DEFAULT 'pending',   -- pending | active | disabled
    failed INTEGER DEFAULT 0, locked_until TEXT,
    created_at TEXT, last_login TEXT);
CREATE TABLE IF NOT EXISTS audit_log(
    id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT, username TEXT,
    action TEXT, detail TEXT);
CREATE TABLE IF NOT EXISTS period_status(period TEXT PRIMARY KEY, locked INTEGER DEFAULT 0);
"""
MAX_FAILS, LOCK_MINUTES = 5, 5
now = lambda: datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def hash_pw(pw: str, salt: bytes | None = None) -> str:
    salt = salt or os.urandom(16)
    h = hashlib.pbkdf2_hmac("sha256", pw.encode(), salt, 200_000)
    return f"pbkdf2${salt.hex()}${h.hex()}"


def verify_pw(pw: str, stored: str) -> bool:
    try:
        _, salt, h = stored.split("$")
        new = hashlib.pbkdf2_hmac("sha256", pw.encode(), bytes.fromhex(salt), 200_000)
        return hmac.compare_digest(new.hex(), h)
    except Exception:
        return False


def audit(username: str, action: str, detail: str = ""):
    with conn() as c:
        c.execute("INSERT INTO audit_log(ts,username,action,detail) VALUES(?,?,?,?)",
                  (now(), username, action, detail))


def user_count() -> int:
    with conn() as c:
        return c.execute("SELECT COUNT(*) FROM users").fetchone()[0]


def _valid(username: str, full_name: str, pw: str):
    if not re.fullmatch(r"[a-z0-9_.]{3,30}", username):
        return "ຊື່ຜູ້ໃຊ້ໃຊ້ໄດ້ສະເພາະ a-z, 0-9, _ ແລະ . (3-30 ຕົວ)"
    if len(full_name.strip()) < 2:
        return "ກະລຸນາໃສ່ຊື່ ແລະ ນາມສະກຸນ"
    if len(pw) < 8:
        return "ລະຫັດຜ່ານຕ້ອງມີຢ່າງໜ້ອຍ 8 ຕົວ"
    return None


def register(username: str, full_name: str, pw: str) -> tuple[bool, str]:
    """ຄົນທຳອິດ = admin (ເປີດໃຊ້ທັນທີ). ຄົນຖັດໄປ = user ລໍຖ້າ admin ອະນຸມັດ."""
    username = username.strip().lower()
    err = _valid(username, full_name, pw)
    if err:
        return False, err
    first = user_count() == 0
    try:
        with conn() as c:
            c.execute(
                "INSERT INTO users(username,full_name,pw_hash,role,status,created_at)"
                " VALUES(?,?,?,?,?,?)",
                (username, full_name.strip(), hash_pw(pw),
                 "admin" if first else "user", "active" if first else "pending", now()))
    except sqlite3.IntegrityError:
        return False, "ຊື່ຜູ້ໃຊ້ນີ້ມີຄົນໃຊ້ແລ້ວ"
    audit(username, "register", "admin ຄົນທຳອິດ" if first else "ລໍຖ້າອະນຸມັດ")
    return True, ("ສ້າງບັນຊີ admin ສຳເລັດ, ເຂົ້າສູ່ລະບົບໄດ້ເລີຍ" if first else
                  "ລົງທະບຽນສຳເລັດ, ກະລຸນາລໍຖ້າ admin ອະນຸມັດ")


def login(username: str, pw: str):
    username = username.strip().lower()
    generic = "ຊື່ຜູ້ໃຊ້ ຫຼື ລະຫັດຜ່ານບໍ່ຖືກຕ້ອງ"
    with conn() as c:
        u = c.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        if not u:
            result = (False, generic, None, None)
        elif u["locked_until"] and u["locked_until"] > now():
            result = (False, f"ບັນຊີຖືກລັອກຊົ່ວຄາວ (ລອງໃໝ່ຫຼັງ {u['locked_until'][11:16]})", None, None)
        elif not verify_pw(pw, u["pw_hash"]):
            f = (u["failed"] or 0) + 1
            lock = (datetime.now() + timedelta(minutes=LOCK_MINUTES)).strftime(
                "%Y-%m-%d %H:%M:%S") if f >= MAX_FAILS else None
            c.execute("UPDATE users SET failed=?, locked_until=? WHERE id=?",
                      (0 if lock else f, lock, u["id"]))
            result = (False, generic, None, "login_fail")
        elif u["status"] == "pending":
            result = (False, "ບັນຊີຍັງລໍຖ້າ admin ອະນຸມັດ", None, None)
        elif u["status"] != "active":
            result = (False, "ບັນຊີນີ້ຖືກປິດການໃຊ້ງານ", None, None)
        else:
            c.execute("UPDATE users SET failed=0, locked_until=NULL, last_login=? WHERE id=?",
                      (now(), u["id"]))
            result = (True, "", dict(id=u["id"], username=u["username"],
                                     full_name=u["full_name"], role=u["role"]), "login")
    ok, msg, user, act = result
    if act:  # ບັນທຶກ audit ຫຼັງປິດ transaction ເພື່ອບໍ່ໃຫ້ຖານຂໍ້ມູນຕິດລັອກ
        audit(username, act)
    return ok, msg, user


def get_user(uid: int):
    with conn() as c:
        r = c.execute("SELECT id,username,full_name,role,status,range_from,range_to,view_all FROM users WHERE id=?",
                      (uid,)).fetchone()
    return dict(r) if r else None


def list_users() -> pd.DataFrame:
    with conn() as c:
        return pd.read_sql_query(
            "SELECT id,username,full_name,role,status,range_from,range_to,view_all,created_at,last_login "
            "FROM users ORDER BY id", c)


def update_user(uid: int, actor: str, status: str | None = None,
                role: str | None = None, new_pw: str | None = None) -> tuple[bool, str]:
    u = get_user(uid)
    if not u:
        return False, "ບໍ່ພົບຜູ້ໃຊ້"
    admins = list_users().query("role=='admin' and status=='active'")
    last_admin = len(admins) == 1 and u["role"] == "admin" and u["status"] == "active"
    if last_admin and ((status and status != "active") or (role and role != "admin")):
        return False, "ບໍ່ສາມາດປິດ/ລົດສິດ admin ຄົນສຸດທ້າຍໄດ້"
    if new_pw is not None and len(new_pw) < 8:
        return False, "ລະຫັດຜ່ານຕ້ອງມີຢ່າງໜ້ອຍ 8 ຕົວ"
    with conn() as c:
        if status:
            c.execute("UPDATE users SET status=? WHERE id=?", (status, uid))
        if role:
            c.execute("UPDATE users SET role=? WHERE id=?", (role, uid))
        if new_pw:
            c.execute("UPDATE users SET pw_hash=?, failed=0, locked_until=NULL WHERE id=?",
                      (hash_pw(new_pw), uid))
    audit(actor, "update_user", f"{u['username']} status={status} role={role} "
          f"pw_reset={bool(new_pw)}")
    return True, "ອັບເດດແລ້ວ"


def change_password(uid: int, old: str, new: str) -> tuple[bool, str]:
    with conn() as c:
        r = c.execute("SELECT pw_hash, username FROM users WHERE id=?", (uid,)).fetchone()
    if not r or not verify_pw(old, r["pw_hash"]):
        return False, "ລະຫັດຜ່ານເກົ່າບໍ່ຖືກຕ້ອງ"
    if len(new) < 8:
        return False, "ລະຫັດຜ່ານໃໝ່ຕ້ອງມີຢ່າງໜ້ອຍ 8 ຕົວ"
    with conn() as c:
        c.execute("UPDATE users SET pw_hash=? WHERE id=?", (hash_pw(new), uid))
    audit(r["username"], "change_password")
    return True, "ປ່ຽນລະຫັດຜ່ານແລ້ວ"


def audit_df(limit: int = 500) -> pd.DataFrame:
    with conn() as c:
        return pd.read_sql_query(
            "SELECT ts,username,action,detail FROM audit_log ORDER BY id DESC LIMIT ?",
            c, params=(limit,))


# ---------- ລັອກງວດ ----------
def period_locked(p: str) -> bool:
    with conn() as c:
        r = c.execute("SELECT locked FROM period_status WHERE period=?", (p,)).fetchone()
    return bool(r and r["locked"])


def set_lock(p: str, locked: bool, actor: str):
    with conn() as c:
        c.execute("INSERT OR REPLACE INTO period_status VALUES(?,?)", (p, int(locked)))
    audit(actor, "lock_period" if locked else "unlock_period", p)


# ---------- ການປ້ອນຂໍ້ມູນຂອງຜູ້ໃຊ້ທົ່ວໄປ (ບວກເພີ່ມເທົ່ານັ້ນ, ແກ້/ລຶບບໍ່ໄດ້) ----------
ENTRY_TYPES = {
    "ຝາກເງິນ": "dep",
    "ຖອນເງິນຝາກ": "wd",
    "ສົ່ງຄືນເງິນກູ້": "repay",
    "ຈ່າຍດອກເບ້ຍເງິນກູ້": "interest_actual",
    "ກູ້ຢືມໃໝ່": "new_loan",
    "ຄ່າໃບຄຳຮ້ອງ": "app_fee",
    "ຄ່າປຶ້ມນ້ອຍ": "booklet_fee",
}


def add_entry(period: str, book_no: int, field: str, amount: float,
              username: str, is_admin: bool = False, note: str = "") -> tuple[bool, str]:
    if field not in ENTRY_TYPES.values():
        return False, "ປະເພດລາຍການບໍ່ຖືກຕ້ອງ"
    if not amount or amount <= 0:
        return False, "ຈຳນວນເງິນຕ້ອງຫຼາຍກວ່າ 0"
    if not is_admin:
        with conn() as c:
            u = c.execute("SELECT range_from, range_to FROM users WHERE username=?",
                          (username,)).fetchone()
        if not u or u["range_from"] is None or u["range_to"] is None:
            return False, "ເຈົ້າຍັງບໍ່ໄດ້ຮັບມອບໝາຍຂອບເຂດເລກທີປຶ້ມ — ກະລຸນາຕິດຕໍ່ Admin"
        rf, rt = u["range_from"], u["range_to"]
        if not (rf <= book_no <= rt):
            audit(username, "entry_denied", f"{period} ປຶ້ມ#{book_no} ນອກຂອບເຂດ {rf}-{rt}")
            return False, f"ເລກທີ {book_no} ບໍ່ຢູ່ໃນຂອບເຂດຮັບຜິດຊອບຂອງເຈົ້າ ({rf}-{rt})"
    if period_locked(period):
        return False, "ເດືອນນີ້ຖືກລັອກແລ້ວ"
    ps = periods()
    if not ps or (not is_admin and period != ps[-1]):
        return False, "ປ້ອນໄດ້ສະເພາະເດືອນຫຼ້າສຸດ"
    with conn() as c:
        r = c.execute("SELECT * FROM monthly WHERE period=? AND book_no=?",
                      (period, book_no)).fetchone()
        if not r:
            return False, "ບໍ່ພົບລູກຄ້າໃນເດືອນນີ້"
        if r["closed"]:
            return False, "ບັນຊີນີ້ຖືກປິດແລ້ວ"
        avail = r["old_dep"] + r["dep"] - r["wd"]
        owed = r["loan_total"] + r["new_loan"] - r["repay"]
        if field == "wd" and amount > avail:
            return False, f"ຖອນເກີນຍອດເງິນຝາກ (ຍອດຖອນໄດ້ {avail:,.0f})"
        if field == "repay" and amount > owed:
            return False, f"ສົ່ງຄືນເກີນຍອດເງິນກູ້ຄ້າງ ({owed:,.0f})"
        old = r[field]
        c.execute(f"UPDATE monthly SET {field}=?, tx_date=? WHERE period=? AND book_no=?",
                  (old + amount, datetime.now().strftime("%Y-%m-%d"), period, book_no))
    audit(username, "entry", f"{period} ປຶ້ມ#{book_no} {field}: {old:,.0f} -> "
          f"{old + amount:,.0f} (+{amount:,.0f}) {note}".strip())
    return True, "ບັນທຶກສຳເລັດ"


# ---------- ຂໍ້ມູນສຳລັບພາບລວມ / ປະຫວັດລູກຄ້າ ----------
def member_history(book_no: int) -> pd.DataFrame:
    with conn() as c:
        df = pd.read_sql_query("SELECT * FROM monthly WHERE book_no=? ORDER BY period",
                               c, params=(book_no,))
    return add_calc(df) if len(df) else df


def trend(rng: tuple[int, int] | None = None) -> pd.DataFrame:
    with conn() as c:
        df = pd.read_sql_query("SELECT * FROM monthly", c)
    if rng:
        df = df[(df.book_no >= rng[0]) & (df.book_no <= rng[1])]
    if df.empty:
        return df
    df = add_calc(df)
    return df.groupby("period")[["dep_next", "loan_next"]].sum()


def backup_bytes() -> bytes:
    return DB_PATH.read_bytes() if DB_PATH.exists() else b""


# ---------- ແບ່ງຂອບເຂດເລກທີປຶ້ມໃຫ້ຜູ້ໃຊ້ແຕ່ລະຄົນ ----------
def compress_ranges(ids: list[int]) -> str:
    """[1,2,3,7,8,10] -> '1-3, 7-8, 10'"""
    ids, out = sorted(ids), []
    i = 0
    while i < len(ids):
        j = i
        while j + 1 < len(ids) and ids[j + 1] == ids[j] + 1:
            j += 1
        out.append(str(ids[i]) if i == j else f"{ids[i]}-{ids[j]}")
        i = j + 1
    return ", ".join(out)


def set_user_scope(uid: int, rf: int, rt: int, view_all: bool, actor: str) -> tuple[bool, str]:
    if rf < 1 or rt < rf:
        return False, "ຂອບເຂດບໍ່ຖືກຕ້ອງ (ເລກທີເລີ່ມຕ້ອງ ≥ 1 ແລະ ບໍ່ເກີນເລກທີສຸດ)"
    u = get_user(uid)
    if not u:
        return False, "ບໍ່ພົບຜູ້ໃຊ້"
    if u["role"] == "admin":
        return False, "Admin ມີສິດທຸກເລກທີຢູ່ແລ້ວ ບໍ່ຕ້ອງກຳນົດຂອບເຂດ"
    others = list_users()
    others = others[(others.id != uid) & (others.role == "user") & (others.status == "active")
                    & others.range_from.notna() & others.range_to.notna()]
    for _, o in others.iterrows():
        if not (rt < o.range_from or rf > o.range_to):
            return False, (f"ຂອບເຂດຊ້ອນກັບ {o.username} ({int(o.range_from)}-{int(o.range_to)}) "
                           "— ປັບຂອບເຂດກ່ອນ ເພື່ອບໍ່ໃຫ້ປ້ອນຊ້ຳຊ້ອນກັນ")
    with conn() as c:
        c.execute("UPDATE users SET range_from=?, range_to=?, view_all=? WHERE id=?",
                  (rf, rt, int(view_all), uid))
    audit(actor, "set_scope", f"{u['username']}: {rf}-{rt} view_all={bool(view_all)}")
    return True, f"ກຳນົດຂອບເຂດ {rf}-{rt} ໃຫ້ {u['username']} ແລ້ວ"


def unassigned_ids() -> list[int]:
    """ເລກທີປຶ້ມທີ່ຍັງບໍ່ມີຜູ້ໃຊ້ (active) ຄົນໃດຮັບຜິດຊອບ"""
    us = list_users()
    us = us[(us.role == "user") & (us.status == "active")
            & us.range_from.notna() & us.range_to.notna()]
    mem = members()
    ids = [int(k) for k in mem.book_no]
    return [k for k in ids
            if not any(r.range_from <= k <= r.range_to for _, r in us.iterrows())]
